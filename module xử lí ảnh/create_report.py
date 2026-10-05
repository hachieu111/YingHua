import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

doc = docx.Document()

# 1. Header
p1 = doc.add_paragraph()
p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
r1 = p1.add_run("TRƯỜNG ĐẠI HỌC LẠC HỒNG\nKHOA CÔNG NGHỆ THÔNG TIN")
r1.bold = True

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r2 = p2.add_run("BÁO CÁO TIẾN ĐỘ TUẦN 08")
r2.bold = True
r2.font.size = Pt(16)

# 2. THÔNG TIN CHUNG
p3 = doc.add_paragraph()
r3 = p3.add_run("THÔNG TIN CHUNG:")
r3.bold = True
r3.underline = True

doc.add_paragraph("- MSSV: 123000831")
doc.add_paragraph("- Họ và tên: Hoàng Hạc Hiếu")
doc.add_paragraph("- Lớp : 23CT111")
doc.add_paragraph("- GVHD: Đoàn Thiện Minh")
doc.add_paragraph("- Tên đề tài: Xây dựng ứng dụng học tiếng Trung thông minh giúp cá nhân hóa người học từ hình ảnh và nội dung câu chuyện")

# 3. THÔNG TIN LÀM VIỆC
p4 = doc.add_paragraph()
r4 = p4.add_run("THÔNG TIN LÀM VIỆC:")
r4.bold = True
r4.underline = True

doc.add_paragraph("Nội dung đã thực hiện trong tuần:")
p_task1 = doc.add_paragraph("- ", style='Normal')
p_task1.add_run("Xây dựng mô-đun tra cứu từ vựng đa chiều (Pinyin):").bold = True
p_task1.add_run(" Tích hợp hiển thị Pinyin đồng thời với chữ Hán gốc, hỗ trợ người học tra cứu cách đọc chuẩn xác trực tiếp trên giao diện tương tác.")

p_task2 = doc.add_paragraph("- ", style='Normal')
p_task2.add_run("Phát triển tính năng tương tác và dịch trực tiếp trên ảnh (Text Rendering):").bold = True
p_task2.add_run(" Hoàn thiện luồng xóa chữ Hán gốc (Inpainting) và đắp bản dịch tiếng Việt lên bong bóng thoại. Triển khai thuật toán Auto-Sizing để căn chỉnh cỡ chữ động dựa trên kích thước chữ Hán gốc, đảm bảo tỷ lệ thẩm mỹ của truyện tranh.")

p_task3 = doc.add_paragraph("- ", style='Normal')
p_task3.add_run("Tối ưu hóa thuật toán nhận diện bong bóng thoại (NMS & Convex Hull):").bold = True
p_task3.add_run(" Xử lý triệt để lỗi mô hình YOLO chia cắt bong bóng thoại bằng cách gộp các vùng Bounding Box sát nhau và bọc đa giác bằng Convex Hull, giúp xóa sạch hoàn toàn chữ Hán.")

p_task4 = doc.add_paragraph("- ", style='Normal')
p_task4.add_run("Xử lý việt hóa Font chữ truyện tranh chuyên dụng:").bold = True
p_task4.add_run(" Sử dụng thư viện fontTools can thiệp trực tiếp vào mã nguồn của bộ font HLcomic, cấy thành công các ký tự bị thiếu (dấu ~, dấu phẩy full-width) để đảm bảo không bị lỗi tàng hình ký tự khi render.")

doc.add_paragraph("Nội dung sẽ thực hiện tiếp trong tuần tiếp theo:")
p_next1 = doc.add_paragraph("- ", style='Normal')
p_next1.add_run("Hoàn thiện mô-đun tra cứu từ vựng đa chiều:").bold = True
p_next1.add_run(" Tiếp tục bổ sung hiển thị chi tiết các trường thông tin: nghĩa tiếng Việt, từ loại và ví dụ minh họa ngữ cảnh cho từng từ vựng.")

p_next2 = doc.add_paragraph("- ", style='Normal')
p_next2.add_run("Phát triển chức năng phân rã chữ Hán:").bold = True
p_next2.add_run(" Tích hợp phân rã cấu trúc nét và bộ thủ để giúp người dùng hiểu sâu thành phần cấu tạo chữ, từ đó tăng khả năng ghi nhớ.")

# 4. Signatures
doc.add_paragraph("\n")
table = doc.add_table(rows=2, cols=2)
table.autofit = True

c00 = table.cell(0, 0)
p_c00 = c00.paragraphs[0]
p_c00.alignment = WD_ALIGN_PARAGRAPH.CENTER
r_c00 = p_c00.add_run("Xác nhận của GVHD")
r_c00.bold = True

c01 = table.cell(0, 1)
p_c01 = c01.paragraphs[0]
p_c01.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_c01.add_run("Đồng Nai, Ngày 27 tháng 9 năm 2026\n").bold = False
r_c01 = p_c01.add_run("Sinh viên")
r_c01.bold = True

c10 = table.cell(1, 0)
p_c10 = c10.paragraphs[0]
p_c10.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_c10.add_run("\n\n\nĐoàn Thiện Minh")

c11 = table.cell(1, 1)
p_c11 = c11.paragraphs[0]
p_c11.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_c11.add_run("\n\n\nHoàng Hạc Hiếu")

# 5. Hình ảnh minh chứng
p_img = doc.add_paragraph("\n\nHình ảnh:")
p_img.runs[0].italic = True

img_path = r"D:\PTUD\test_tilde.png"
if os.path.exists(img_path):
    doc.add_picture(img_path, width=Inches(5.0))

doc.save(r"d:\PTUD\123000831_HoangHacHieu_Tuan08_DoanThienMinh.docx")
print("Đã tạo file Word báo cáo thành công!")
